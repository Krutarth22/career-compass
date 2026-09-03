"""Unit tests for markdown_docx.py."""

from markdown_docx import render_markdown_to_docx


def _paragraph_texts(doc):
    return [p.text for p in doc.paragraphs]


def test_headers_produce_correct_heading_levels():
    doc = render_markdown_to_docx("# H1\n## H2\n### H3\n#### H4\n")

    styles = [(p.style.name, p.text) for p in doc.paragraphs]
    assert styles == [
        ("Heading 1", "H1"),
        ("Heading 2", "H2"),
        ("Heading 3", "H3"),
        ("Heading 4", "H4"),
    ]


def test_bold_and_italic_spans_produce_formatted_runs():
    doc = render_markdown_to_docx("This is **bold** and *italic* text.")

    para = doc.paragraphs[0]
    bold_runs = [r for r in para.runs if r.bold]
    italic_runs = [r for r in para.runs if r.italic]

    assert any(r.text == "bold" for r in bold_runs)
    assert any(r.text == "italic" for r in italic_runs)
    assert "This is" in para.text
    assert "bold" in para.text
    assert "italic" in para.text


def test_bullet_list_produces_list_bullet_paragraphs():
    doc = render_markdown_to_docx("- item one\n- item two\n- item three\n")

    bullet_paras = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert [p.text for p in bullet_paras] == ["item one", "item two", "item three"]


def test_numbered_list_produces_list_number_paragraphs():
    doc = render_markdown_to_docx("1. step one\n2. step two\n")

    number_paras = [p for p in doc.paragraphs if p.style.name == "List Number"]
    assert [p.text for p in number_paras] == ["step one", "step two"]


def test_pipe_table_produces_table_with_correct_rows_and_cells():
    md = "| Skill | Tier | Status |\n" "|---|---|---|\n" "| Python | Critical | open |\n" "| SQL | High | practiced |\n"

    doc = render_markdown_to_docx(md)

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert [c.text for c in table.rows[0].cells] == ["Skill", "Tier", "Status"]
    assert [c.text for c in table.rows[1].cells] == ["Python", "Critical", "open"]
    assert [c.text for c in table.rows[2].cells] == ["SQL", "High", "practiced"]


def test_table_header_row_is_bolded():
    md = "| Skill | Tier |\n|---|---|\n| Python | Critical |\n"

    doc = render_markdown_to_docx(md)

    header_cell_para = doc.tables[0].rows[0].cells[0].paragraphs[0]
    assert all(run.bold for run in header_cell_para.runs)


def test_markdown_link_produces_hyperlink_run():
    doc = render_markdown_to_docx("See [the docs](https://example.com/docs) for more.")

    para = doc.paragraphs[0]
    hyperlink_els = para._p.findall(
        ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink"
    )
    assert len(hyperlink_els) == 1

    rels = para.part.rels
    hyperlink_rels = [r for r in rels.values() if r.reltype.endswith("hyperlink")]
    assert any(r.target_ref == "https://example.com/docs" for r in hyperlink_rels)


def test_paragraphs_separated_by_blank_lines_stay_separate():
    doc = render_markdown_to_docx("First paragraph.\n\nSecond paragraph.\n")

    texts = [p.text for p in doc.paragraphs if p.text]
    assert texts == ["First paragraph.", "Second paragraph."]


def test_unrecognized_construct_does_not_raise():
    # A stray horizontal rule and an unusual line shouldn't crash the
    # renderer -- it should degrade gracefully.
    doc = render_markdown_to_docx("Some text.\n\n---\n\nMore text.\n")

    texts = [p.text for p in doc.paragraphs if p.text]
    assert "Some text." in texts
    assert "More text." in texts


def test_render_into_existing_document_appends():
    import docx

    doc = docx.Document()
    doc.add_paragraph("Preexisting content")

    render_markdown_to_docx("# New Section\nNew body.", doc=doc)

    texts = _paragraph_texts(doc)
    assert texts[0] == "Preexisting content"
    assert "New Section" in texts
    assert "New body." in texts
